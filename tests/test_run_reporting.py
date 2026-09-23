"""Run isolation, failure injection, reporting integrity, and attribution guards."""
import json
from pathlib import Path
import shutil
import socket
import unittest
from unittest.mock import patch
import uuid

import pandas as pd
from openpyxl import load_workbook

from pipeline.config import Config
from pipeline.pipeline_v053 import run_pipeline, read_dataset_files
from pipeline.run_reporting import RunSession, atomic_json, file_hash, recover_interrupted_runs
from pipeline.rasch_quality import fit_rasch_quality
from pipeline.qc_pipeline import reconstruct_qc
from pipeline.data_contracts import item_mapping


class RunReportingTests(unittest.TestCase):
    def setUp(self):
        self.path = Path(__file__).parent / "_artifacts" / ("report_" + uuid.uuid4().hex)
        self.path.mkdir(parents=True)
        self.addCleanup(shutil.rmtree, self.path)
        self.inputs = self.path / "input"
        (self.inputs / "01_Production").mkdir(parents=True)
        (self.inputs / "02_Planner_Skills").mkdir()
        rows = [{"Reference": f"WO{i}", "RoundNo": 1, "Worker": "A" if i % 2 else "B",
                 "Item Number": "001" if i % 3 else "002", "Process": "Sanding",
                 "Qty Doing": 10, "Total Actual Hours": 1 + i * .1, "RAF Month": f"2026-08-{i+1:02d}"}
                for i in range(12)]
        pd.DataFrame(rows).to_csv(self.inputs / "01_Production" / "p.csv", index=False)
        pd.DataFrame({"Worker ID": ["A", "B"], "Process": ["Sanding"] * 2,
                      "Planner Verified Skill Level": [5, 6]}).to_csv(self.inputs / "02_Planner_Skills" / "p.csv", index=False)
        self.config = Config(input_dir=self.inputs, output_dir=self.path / "out")

    def last_summary(self):
        paths = sorted(self.config.output_dir.glob("run_*/run_summary.json"))
        return json.loads(paths[-1].read_text(encoding="utf-8"))

    def test_success_repeatability_and_manifest_integrity(self):
        first = run_pipeline(self.config)
        second = run_pipeline(self.config)
        self.assertNotEqual(first["output_dir"], second["output_dir"])
        self.assertEqual(first["summary"]["data"], second["summary"]["data"])
        self.assertEqual(second["summary"]["comparison"]["status"], "COMPARABLE")
        self.assertIsNone(second["summary"]["qc"]["eligible_rounds"])
        self.assertEqual(second["summary"]["decision_status"], "DIAGNOSTIC_ONLY")
        for entry in second["manifest"]["artifacts"]:
            self.assertEqual(file_hash(second["output_dir"] / entry["path"]), entry["sha256"])
        workbook = load_workbook(second["output_dir"] / "artifacts" / "Skill_ID_Ket_qua_chay_thu.xlsx", read_only=True)
        try:
            self.assertEqual(workbook.sheetnames[0], "Run Summary")
            rows = dict(workbook["Run Summary"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(rows["Execution Status"], second["summary"]["execution_status"])
        finally:
            workbook.close()

    def test_atomic_json_retries_transient_lock_and_rejects_persistent_lock(self):
        target = self.path / "checkpoint.json"
        atomic_json(target, {"value": 1})
        original = Path.replace
        calls = []
        def transient(source, destination):
            calls.append(source)
            if len(calls) == 1:
                raise PermissionError("temporary destination lock")
            return original(source, destination)
        with patch.object(Path, "replace", transient), patch("pipeline.run_reporting.time.sleep"):
            atomic_json(target, {"value": 2})
        self.assertEqual(json.loads(target.read_text())["value"], 2)
        with patch.object(Path, "replace", side_effect=PermissionError("persistent lock")), patch("pipeline.run_reporting.time.sleep"):
            with self.assertRaises(PermissionError):
                atomic_json(target, {"value": 3})
        self.assertEqual(json.loads(target.read_text())["value"], 2)
        self.assertEqual(list(self.path.glob("*.tmp")), [])

    def test_failed_export_preserves_successful_pointer(self):
        run_pipeline(self.config)
        pointer = self.config.output_dir / "latest_successful.json"
        before = pointer.read_bytes()
        with patch("pipeline.workbook_io.write_result_workbook", side_effect=PermissionError("locked workbook")):
            with self.assertRaises(PermissionError):
                run_pipeline(self.config)
        self.assertEqual(pointer.read_bytes(), before)
        summary = self.last_summary()
        self.assertEqual(summary["execution_status"], "FAILED")
        self.assertEqual(summary["failed_stage"], "Artifact export")
        self.assertTrue(summary["artifacts"])

    def test_missing_and_corrupt_input_have_failure_reports(self):
        config = Config(input_dir=self.path / "missing", output_dir=self.config.output_dir)
        with self.assertRaises(FileNotFoundError):
            run_pipeline(config)
        self.assertEqual(self.last_summary()["failed_stage"], "Input snapshot")
        (self.inputs / "01_Production" / "bad.xlsx").write_text("not a workbook")
        with self.assertRaises(ValueError):
            run_pipeline(self.config)
        self.assertEqual(self.last_summary()["failed_stage"], "Input loading")
        self.assertFalse((self.config.output_dir / "latest_successful.json").exists())

    def test_snapshot_mutation_is_rejected(self):
        original_copy = shutil.copyfile
        def mutate(source, target):
            result = original_copy(source, target)
            with Path(source).open("a") as stream:
                stream.write("\n")
            return result
        with patch("pipeline.run_reporting.shutil.copyfile", side_effect=mutate):
            with self.assertRaisesRegex(ValueError, "Input changed"):
                run_pipeline(self.config)
        self.assertEqual(self.last_summary()["execution_status"], "FAILED")

    def test_keyboard_interrupt_is_recorded(self):
        with patch("pipeline.pipeline_v053.fit_time_models", side_effect=KeyboardInterrupt):
            with self.assertRaises(KeyboardInterrupt):
                run_pipeline(self.config)
        summary = self.last_summary()
        self.assertEqual(summary["execution_status"], "INTERRUPTED")
        self.assertEqual(summary["exit_code"], 130)

    def test_model_exception_and_report_finalization_failure(self):
        with patch("pipeline.pipeline_v053.fit_time_models", side_effect=RuntimeError("fit failed")):
            with self.assertRaises(RuntimeError):
                run_pipeline(self.config)
        self.assertEqual(self.last_summary()["failed_stage"], "Time models")
        with patch("pipeline.run_reporting.RunSession.report_rows", side_effect=RuntimeError("summary sheet failed")):
            with self.assertRaises(RuntimeError):
                run_pipeline(self.config)
        self.assertEqual(self.last_summary()["execution_status"], "FAILED")
        self.assertFalse((self.config.output_dir / "latest_successful.json").exists())

    def test_cli_and_api_defaults_match(self):
        from pipeline.main import main
        with patch("sys.argv", ["pipeline.main"]), patch("pipeline.main.run_pipeline", return_value={"output_dir": self.path}) as runner:
            main()
        config = runner.call_args.args[0]
        self.assertEqual(config.incomplete_month, Config().incomplete_month)
        self.assertEqual(config.qc_exclude_month, Config().qc_exclude_month)

    def test_nonconverged_first_pass_is_not_labelled_calibrated(self):
        (self.inputs / "04_QC_Tickets").mkdir()
        production = pd.read_csv(self.inputs / "01_Production" / "p.csv", dtype={"Item Number": str})
        qc = []
        for i, row in production.iterrows():
            for status, qty in (("Pass", 6), ("Fail", 4)):
                qc.append({"QualityOrderId": f"q{i}{status}", "WO": row.Reference, "Item Number": row["Item Number"],
                           "RoundNo": 1, "QCStatus": status, "QCQty": qty, "ExpectedQty": 10,
                           "QC_Start": row["RAF Month"], "QC_Stop": row["RAF Month"]})
        pd.DataFrame(qc).to_csv(self.inputs / "04_QC_Tickets" / "q.csv", index=False)
        config = Config(input_dir=self.inputs, output_dir=self.config.output_dir, rasch_max_iter=1)
        result = run_pipeline(config)
        items = result["tables"]["Do kho SKU"]
        self.assertTrue(items.FirstPass_Fit_Status.eq("NONCONVERGED").all())
        self.assertTrue(items.FirstPass_Decision_Eligibility.eq("BLOCKED").all())
        self.assertFalse(items["Quality Status"].str.contains("calibrated").any())
        self.assertTrue(items["QualityDifficulty"].isna().all())

    def test_recover_abandoned_but_not_active_run(self):
        session = RunSession(self.config)
        session.start_stage("Input loading")
        session.close()
        with patch("pipeline.run_reporting.process_alive", return_value=True):
            self.assertEqual(recover_interrupted_runs(self.config.output_dir), [])
        with patch("pipeline.run_reporting.process_alive", return_value=False):
            self.assertEqual(recover_interrupted_runs(self.config.output_dir), [session.run_id])
        self.assertEqual(self.last_summary()["execution_status"], "INTERRUPTED")
        self.assertFalse((self.config.output_dir / "latest_successful.json").exists())

    def test_changed_cohort_not_directly_comparable(self):
        run_pipeline(self.config)
        path = self.inputs / "01_Production" / "p.csv"
        frame = pd.read_csv(path, dtype={"Item Number": str})
        frame.loc[0, "Total Actual Hours"] = 2
        frame.to_csv(path, index=False)
        result = run_pipeline(self.config)
        self.assertEqual(result["summary"]["comparison"]["status"], "NOT_DIRECTLY_COMPARABLE")
        self.assertNotIn("wape_ratio_changes", result["summary"]["comparison"])

    def test_missing_rounds_stay_unknown(self):
        path = self.inputs / "01_Production" / "p.csv"
        pd.read_csv(path, dtype={"Item Number": str}).drop(columns="RoundNo").to_csv(path, index=False)
        result = run_pipeline(self.config)
        self.assertTrue(result["data"].RoundNo.isna().all())
        self.assertEqual(result["summary"]["data"]["unknown_round_rows"], 12)

    def test_master_duplicates_do_not_multiply_rows(self):
        path = self.path / "mes.xlsx"
        with pd.ExcelWriter(path) as writer:
            pd.DataFrame({"ProductionOrderNumber": ["W"], "Worker": ["A"], "Final": [1], "RoundNo": [1]}).to_excel(writer, sheet_name="GSWorkerWorkingHours", index=False)
            pd.DataFrame({"ProductionOrderNumber": ["W", "W"], "GoodCW": [1, 2]}).to_excel(writer, sheet_name="WorkOrderData", index=False)
        with self.assertRaisesRegex(ValueError, "duplicate ProductionOrderNumber"):
            read_dataset_files([path], "Production")
        # Reader must release the file even after rejection.
        path.rename(self.path / "renamed.xlsx")

    def test_unknown_worker_never_falls_back_to_primary(self):
        first = pd.DataFrame({"WO": ["WO1", "WO2"], "Worker": [None, None], "Process": ["Sanding"] * 2,
                              "SizeAdjustedGroup": ["G1", "G2"], "PassQty": [5, 5], "InspectedQty": [10, 10]})
        production = pd.DataFrame({"Reference": ["WO1", "WO2"], "Worker": ["A", "B"], "RoundNo": [2, 2], "Total Actual Hours": [4, 5]})
        result = fit_rasch_quality(first, production)
        self.assertTrue(result["group_difficulty"].empty)

    def test_nonconverged_diagnostics_and_all_fail_fit_are_finite(self):
        first = pd.DataFrame({"WO": [f"W{i}" for i in range(8)], "Worker": ["A", "B"] * 4,
                              "Process": ["Sanding"] * 8, "SizeAdjustedGroup": ["G1"] * 4 + ["G2"] * 4,
                              "PassQty": [0] * 8, "InspectedQty": [10] * 8})
        production = first[["WO", "Worker"]].rename(columns={"WO": "Reference"})
        groups = fit_rasch_quality(first, production, max_iter=1)["group_difficulty"]
        self.assertTrue(groups.FPY_Rasch_Difficulty.notna().all())
        self.assertTrue(groups.DecisionEligibility.eq("BLOCKED").all())
        self.assertTrue(groups.CalibrationStatus.eq("NOT_VALIDATED").all())

    def test_qc_exclusion_uses_created_date_and_can_be_disabled(self):
        mapping = item_mapping(pd.DataFrame({"Item Number": ["001"], "Process": ["Sanding"]}))
        tickets = pd.DataFrame({"QualityOrderId": ["Q1"], "WO": ["W1"], "Item Number": ["001"], "RoundNo": [1],
                                "QCStatus": ["Pass"], "QCQty": [1], "ExpectedQty": [1],
                                "CreatedDateTime": ["2026-07-31"], "QC_Start": ["2026-08-01"], "QC_Stop": ["2026-08-01"]})
        rounds, _, errors = reconstruct_qc(tickets, mapping)
        self.assertTrue(rounds.empty)
        self.assertEqual(len(errors), 1)
        rounds, _, _ = reconstruct_qc(tickets, mapping, exclude_month="")
        self.assertEqual(len(rounds), 1)
        self.assertEqual(Config().incomplete_month, "")


if __name__ == "__main__":
    unittest.main()
