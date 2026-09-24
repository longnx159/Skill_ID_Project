"""Business-rule regression tests. All records here are synthetic fixtures."""
import tempfile
import uuid
from contextlib import nullcontext
import unittest
from pathlib import Path
import numpy as np
import pandas as pd
from pipeline.config import Config
from pipeline.preprocessing import clean_production_data, clean_planner_data
from pipeline.data_contracts import item_mapping, attach_mapping, chronological_split, connectivity, GROUP
from pipeline.qc_pipeline import reconstruct_qc, recovery_trajectories, recovery_groups, reconstruct_touch
from pipeline.scoring import technical_complexity, FACTOR_WEIGHTS, confidence_tier
from pipeline.hybrid_effect import fit_hybrid_effect


def mapping():
    return item_mapping(pd.DataFrame({"Item Number":["001","002"],"Process":["Sanding"]*2,"Item Number (Size Adjusted)":["G","G"]}))


def tickets(wo="WO1", day="2026-08-01"):
    return pd.DataFrame([
        ["q1",wo,"001",1,"Pass",6,10,day+" 10:00",day+" 10:10","A"],
        ["q2",wo,"001",1,"Fail",4,10,day+" 10:00",day+" 10:10","A"],
        ["q3",wo,"001",2,"Pass",4,4,"2026-08-02 10:00","2026-08-02 10:10","A"],
    ], columns=["QualityOrderId","WO","Item Number","RoundNo","QCStatus","QCQty","ExpectedQty","QC_Start","QC_Stop","InitialWorker"])


class Rules(unittest.TestCase):
    def test_fractional_qty_and_latest_month(self):
        raw=pd.DataFrame({"Worker":["A","A"],"Item Number":["001"]*2,"Qty Doing":["0.5","1"],"Total Actual Hours":["1","1"],"RAF Month":["2026-06-01","2026-07-01"]})
        out=clean_production_data(raw)
        self.assertEqual(len(out),2)
        self.assertEqual(out["Minutes per final OK"].iloc[0],120)
        self.assertEqual(len(clean_production_data(raw,"2026-07")),1)

    def test_mapping_conflicts_and_fallback(self):
        raw=pd.DataFrame({"Item Number":["A","A"],"Process":["P","P"],"Item Number (Size Adjusted)":["X","Y"]})
        with self.assertRaises(ValueError): item_mapping(raw)
        out=item_mapping(raw.drop(columns="Item Number (Size Adjusted)"))
        self.assertEqual(out[GROUP].iloc[0],"A")
        self.assertTrue(out.MappingFallback.all())
        bad=pd.DataFrame({"Item Number":["001"],"Process":["Soldering"]})
        with self.assertRaises(ValueError): attach_mapping(bad,mapping())

    def test_planner_raw_skill_and_duplicates(self):
        p=pd.DataFrame({"Worker ID":["A"],"Process":["P"],"Planner Verified Skill Level":[1.43]})
        self.assertEqual(clean_planner_data(p).iloc[0]["Planner Verified Skill Level"],1.43)
        with self.assertRaises(ValueError): clean_planner_data(pd.concat([p,p]))

    def test_qc_rounds_and_denominator(self):
        rounds,first,errors=reconstruct_qc(tickets(),mapping())
        self.assertTrue(errors.empty)
        self.assertEqual(first.FPY.iloc[0],.6)
        self.assertEqual(first.InspectedQty.iloc[0],10)
        self.assertEqual(first.Worker.iloc[0],"A")
        self.assertEqual(len(rounds),2)

    def test_july_excludes_rows_but_keeps_later_wo_rounds(self):
        q=tickets()
        july=q.iloc[[0]].copy()
        july["QualityOrderId"]="july"
        july["QC_Start"]="2026-07-31 10:00"
        july["QC_Stop"]="2026-07-31 10:10"
        rounds,first,errors=reconstruct_qc(pd.concat([q,july]),mapping())
        self.assertFalse(rounds.empty)
        self.assertTrue(pd.to_datetime(rounds.QC_Start).ge("2026-08-01").all())
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors.Reason.iloc[0], "July QC excluded")

    def test_duplicate_and_reconciliation(self):
        q=tickets()
        for change in ["duplicate","overflow","roundgap"]:
            bad=q.copy()
            if change=="duplicate": bad.loc[2,"QualityOrderId"]="q1"
            elif change=="overflow": bad.loc[2,["QCQty","ExpectedQty"]]=[5,5]
            else: bad.loc[2,"RoundNo"]=3
            rounds,_,errors=reconstruct_qc(bad,mapping())
            self.assertTrue(rounds.empty,change)
            self.assertGreater(len(errors),0)

    def test_cutoff_does_not_use_future_recovery(self):
        rounds,_,_=reconstruct_qc(tickets(),mapping(),"2026-08-01 23:59:59")
        r=recovery_trajectories(rounds,"2026-08-01 23:59:59")
        self.assertTrue(r.OpenWO.iloc[0])
        self.assertTrue(np.isnan(r.ObservedRecoveryDifficulty.iloc[0]))
        self.assertEqual(r.RecoverySensitivityRange.iloc[0],10)
        self.assertTrue(r.ScrapQty.isna().all())
        g=recovery_groups(r,open_threshold=.2,sensitivity_threshold=2,approvals=True)
        self.assertEqual(g.RecoveryStabilityFlag.iloc[0],"RecoveryUnstable")

    def test_recovery_zero_fail_is_na_and_decay(self):
        rounds,_,_=reconstruct_qc(tickets(),mapping())
        self.assertEqual(recovery_trajectories(rounds,"2026-08-03").ObservedRecoveryDifficulty.iloc[0],0)
        zero=rounds.iloc[[0]].copy()
        zero["FailQty"]=0
        self.assertTrue(np.isnan(recovery_trajectories(zero,"2026-08-03").ObservedRecoveryDifficulty.iloc[0]))
        third=rounds.copy()
        third.loc[third.RoundNo.eq(2),"RoundNo"]=3
        self.assertAlmostEqual(recovery_trajectories(third,"2026-08-03").ObservedRecoveryDifficulty.iloc[0],3.5)

    def test_touch_idle_process_guard_and_rescue(self):
        rounds,_,_=reconstruct_qc(tickets(),mapping())
        events=pd.DataFrame([
            ["WO1","A","Sanding","2026-08-01 09:00","2026-08-01 09:30"],
            ["WO1","A","Sanding","2026-08-02 09:00","2026-08-02 09:15"],
            ["WO1","B","Sanding","2026-08-02 09:20","2026-08-02 09:30"],
            ["WO1","B","Polishing","2026-08-02 09:40","2026-08-02 09:50"],
            ["WO1","B","Sanding","2026-08-02 09:50","2026-08-02 09:40"],
        ],columns=["WO","Worker","Process","Worker_Start","Worker_Stop"])
        out=reconstruct_touch(events,rounds)
        self.assertEqual(out.Phase.tolist(),["FirstPass","Rework","Rework","OtherProcess","InvalidTouch"])
        self.assertEqual(out.Attribution.iloc[1],"SelfRework")
        self.assertEqual(out.Attribution.iloc[2],"AssistedRescue")
        self.assertEqual(out.TouchMinutes.sum(),55)

    def test_wo_temporal_embargo(self):
        d=pd.DataFrame({"WO":["A","B","B","C","D"],"Date":pd.to_datetime(["2026-01-01","2026-02-01","2026-05-01","2026-04-01","2026-05-01"])})
        train,test,embargo=chronological_split(d,"Date","WO",.5)
        self.assertFalse(set(train.WO)&set(test.WO))
        self.assertEqual(set(embargo.WO),{"B"})
        self.assertLess(train.Date.max(),test.Date.min())

    def test_network_disconnected_not_dropped(self):
        d=pd.DataFrame({"Process":["P"]*4,"Worker":["A","B","C","C"],GROUP:["G1","G1","G2","G2"],"CertifiedSkill":[5,6,9,9]})
        n=connectivity(d,min_n=2)
        self.assertEqual(n.ConnectedComponentID.nunique(),2)
        self.assertTrue(n.loc[n[GROUP].eq("G2"),"ExtrapolationFlag"].iloc[0])

    def test_factors_missing_not_zero(self):
        f=pd.DataFrame([{"Item Number":"I","Process":"P","Factor":factor,"Score":5,"Evidence":"E","Scorer":"S","Approver":"A","Version":"v1","Approved":True,"DataQualityStatus":"Valid","Method":"MODEL_ESTIMATE" if factor=="Quality" else "FIXED_INPUT","ConfidencePct":0.8 if factor=="Quality" else np.nan,"ConfidenceStatus":"ESTIMATED_VALIDATED" if factor=="Quality" else "NOT_APPLICABLE"} for factor in FACTOR_WEIGHTS])
        self.assertEqual(technical_complexity(f)["Final Technical Complexity"].iloc[0],5)
        f.loc[0,"Score"]=np.nan
        self.assertTrue(np.isnan(technical_complexity(f)["Final Technical Complexity"].iloc[0]))
        self.assertEqual(confidence_tier(3,5)[0],"Low")

    def test_modeled_factor_needs_validated_confidence(self):
        f=pd.DataFrame([{"Item Number":"I","Process":"P","Factor":factor,"Score":5,"Evidence":"E","Scorer":"S","Approver":"A","Version":"v1","Approved":True,"DataQualityStatus":"Valid","Method":"MODEL_ESTIMATE" if factor in {"Quality","Learning"} else "FIXED_INPUT","ConfidencePct":0.8 if factor=="Quality" else np.nan,"ConfidenceStatus":"ESTIMATED_VALIDATED" if factor=="Quality" else "NOT_ESTIMATED"} for factor in FACTOR_WEIGHTS])
        result=technical_complexity(f).iloc[0]
        self.assertTrue(np.isnan(result["Final Technical Complexity"]))
        self.assertIn("Model confidence missing or unvalidated",result.FactorGate)
        f.loc[f.Factor.eq("Learning"),["ConfidencePct","ConfidenceStatus"]]=[0.7,"ESTIMATED_VALIDATED"]
        self.assertEqual(technical_complexity(f)["Final Technical Complexity"].iloc[0],5)
        f.loc[f.Factor.eq("Quality"),"ConfidenceStatus"]="CONDITIONAL_DIAGNOSTIC"
        self.assertTrue(np.isnan(technical_complexity(f)["Final Technical Complexity"].iloc[0]))

    def test_one_worker_model_finite(self):
        d=pd.DataFrame({"Worker":["A"]*3,"Item Number":["I"]*3,"log_time":[1.,1.,1.]})
        m=fit_hybrid_effect(d)
        self.assertTrue(np.isfinite(m["predictions"]).all())

    def test_time_em_likelihood_matches_dense_covariance(self):
        d=pd.DataFrame({"Worker":["B","A","B","A","B","A"],"Item Number":["I","I","J","J","I","J"],"log_time":[.5,-.4,1.1,.2,.6,.1],"Qty Doing":[1,2,3,4,2,1]})
        m=fit_hybrid_effect(d)
        variance=m["variance_components"].set_index("Component").Variance
        w=np.sqrt(d["Qty Doing"].to_numpy(float)); w/=w.mean()
        z=pd.get_dummies(d.Worker).to_numpy(float)
        cov=np.diag(variance["Residual"]/w)+variance["Worker random effect"]*(z@z.T)
        effects=m["item_scores"].set_index("Item Number")["Hybrid Item FE"]
        residual=d.log_time.to_numpy()-m["intercept"]-d["Item Number"].map(effects).to_numpy()
        expected=-.5*(len(d)*np.log(2*np.pi)+np.linalg.slogdet(cov)[1]+residual@np.linalg.solve(cov,residual))
        self.assertAlmostEqual(expected,m["log_likelihood"],places=7)
        self.assertTrue(m["converged"])

    def test_quality_scaling_and_frozen_prediction(self):
        from pipeline.quality_model import SkillScaler,predict_quality,quality_metrics,attach_asof_skill
        scaler=SkillScaler.fit(pd.Series([4.,5.,6.]))
        self.assertEqual(scaler.transform([5])[0],0)
        self.assertEqual(scaler.sd,1)
        posterior={"alpha":np.zeros(20),"beta":np.ones(20),"u":np.zeros((20,1)),"v":np.zeros((20,1)),"worker_sd":np.zeros(20),"group_sd":np.zeros(20)}
        model={"posterior":posterior,"workers":["A"],"groups":["G"],"scaler":scaler,"training_wos":frozenset(["train"]),"likelihood":"beta_binomial"}
        test=pd.DataFrame({"WO":["test"],"Worker":["B"],GROUP:["H"],"Process":["P"],"PassQty":[5],"InspectedQty":[10],"CertifiedSkill":[5]})
        pred=predict_quality(model,test)
        self.assertEqual(pred.PredictedFPY.iloc[0],.5)
        self.assertTrue(pred.UnseenWorker.iloc[0] and pred.UnseenGroup.iloc[0])
        self.assertAlmostEqual(quality_metrics(pred)["Piece Log Loss"],np.log(2))
        test["WO"]="train"
        with self.assertRaises(ValueError): predict_quality(model,test)
        first=pd.DataFrame({"Worker":["A"],"Process":["P"],"QC_Start":pd.to_datetime(["2026-08-01"])})
        planner=pd.DataFrame({"Worker ID":["A"],"Process":["P"],"Planner Verified Skill Level":[9],"EffectiveFrom":["2026-09-01"]})
        self.assertTrue(attach_asof_skill(first,planner,allow_baseline_prior=False).CertifiedSkill.isna().all())
        self.assertEqual(attach_asof_skill(first,planner,allow_baseline_prior=True).CertifiedSkill.iloc[0], 9)

    def test_template_and_pipeline_end_to_end(self):
        from pipeline.create_template import create_template
        from openpyxl import load_workbook
        from pipeline.pipeline_v053 import run_pipeline
        # Windows sandbox ACLs reject tempfile's mode-0700 directories.
        fixture_dir=Path(__file__).resolve().parent/"_artifacts"/uuid.uuid4().hex
        fixture_dir.mkdir(parents=True)
        with nullcontext(fixture_dir) as tmp:
            path=Path(tmp)/"input.xlsx"
            create_template(path)
            wb=load_workbook(path)
            def populate(sheet, records):
                headers=[c.value for c in wb[sheet][1]]
                extra=[k for k in records[0].keys() if k not in headers]
                for i, k in enumerate(extra, len(headers)+1): wb[sheet].cell(1, i, k)
                headers += extra
                for r,record in enumerate(records,2):
                    for c,h in enumerate(headers,1): wb[sheet].cell(r,c,record.get(h))
            populate("Production",[{"Reference":f"WO{i}","Worker":"A" if i%2 else "B","Item Number":"001","Item Number (Size Adjusted)":"G","Process":"Sanding","Qty Doing":10,"Total Actual Hours":1+i*.01,"RAF Month":pd.Timestamp(f"2026-0{1+i%7}-01").to_pydatetime()} for i in range(35)])
            populate("Planner Skills",[{"Worker ID":"A","Process":"Sanding","Planner Verified Skill Level":5},{"Worker ID":"B","Process":"Sanding","Planner Verified Skill Level":6}])
            populate("QC Tickets",tickets().to_dict("records"))
            wb.save(path)
            result=run_pipeline(Config(output_dir=Path(tmp)/"out", incomplete_month=""),path)
            self.assertEqual(len(result["data"]),35)
            items=result["tables"]["Do kho SKU"]
            self.assertEqual(items.FPY_Raw_Difficulty.iloc[0],4)
            self.assertTrue(items["Final Technical Complexity"].isna().all())
            self.assertTrue((result["output_dir"]/"artifacts"/"Skill_ID_Ket_qua_chay_thu.xlsx").exists())

    def test_incremental_folder_ingestion_combines_files_and_skips_lock(self):
        from pipeline.pipeline_v053 import load_inputs, discover_input_files, run_pipeline
        from pipeline.config import Config
        fixture_dir=Path(__file__).resolve().parent/"_artifacts"/uuid.uuid4().hex
        (fixture_dir/"01_Production").mkdir(parents=True)
        (fixture_dir/"02_Planner_Skills").mkdir(parents=True)
        production=pd.DataFrame({"Reference":["WO1"],"Worker":["A"],"Item Number":["001"],"Process":["Sanding"],"Qty Doing":[1],"Total Actual Hours":[1],"RAF Month":["2026-08-01"]})
        production.to_csv(fixture_dir/"01_Production"/"production_01.csv",index=False)
        production.assign(Reference="WO2").to_csv(fixture_dir/"01_Production"/"production_02.csv",index=False)
        production.to_csv(fixture_dir/"01_Production"/"~$locked.csv",index=False)
        planner=pd.DataFrame({"Worker ID":["A"],"Process":["Sanding"],"Planner Verified Skill Level":[5]})
        planner.to_csv(fixture_dir/"02_Planner_Skills"/"planner.csv",index=False)
        sources=load_inputs(Config(input_dir=fixture_dir))
        self.assertEqual(len(discover_input_files(fixture_dir)["Production"]),2)
        self.assertEqual(len(sources["Production"]),2)
        self.assertEqual(set(sources["Production"].Reference),{"WO1","WO2"})
        self.assertEqual(sources["Production"].SourceFile.nunique(),2)
        result=run_pipeline(Config(input_dir=fixture_dir,output_dir=fixture_dir/"out"))
        self.assertEqual(len(result["data"]),2)
        self.assertTrue((result["output_dir"]/"artifacts"/"Skill_ID_Ket_qua_chay_thu.xlsx").exists())


    def test_mes_production_and_qc_database_ingestion(self):
        from pipeline.pipeline_v053 import run_pipeline
        from pipeline.config import Config
        fixture_dir = Path(__file__).resolve().parent / "_artifacts" / uuid.uuid4().hex
        (fixture_dir / "01_Production").mkdir(parents=True)
        (fixture_dir / "02_Planner_Skills").mkdir(parents=True)
        (fixture_dir / "04_QC_Tickets").mkdir(parents=True)
        (fixture_dir / "09_Item_Master").mkdir(parents=True)

        # Multi-sheet MES ProductionData
        with pd.ExcelWriter(fixture_dir / "01_Production" / "ProductionData.xlsx") as writer:
            worker_hours = pd.DataFrame([
                {"Worker": "W1", "Name": "Worker 1", "Department": "SAN 1", "ProductionOrderNumber": "WO100", "RoundNo": 1, "Final": 3.0},
                {"Worker": "W2", "Name": "Worker 2", "Department": "SAN 1", "ProductionOrderNumber": "WO100", "RoundNo": 1, "Final": 1.0},
                {"Worker": "W1", "Name": "Worker 1", "Department": "SAN 1", "ProductionOrderNumber": "WO101", "RoundNo": 1, "Final": 2.0},
            ])
            wo_data = pd.DataFrame([
                {"ProductionOrderNumber": "WO100", "Status": "Complete", "MaxRAFDate": "2026-08-01", "ItemNumber": "2SN001", "GoodCW": 20},
                {"ProductionOrderNumber": "WO101", "Status": "Complete", "MaxRAFDate": "2026-08-01", "ItemNumber": "2SN001", "GoodCW": 10},
            ])
            worker_hours.to_excel(writer, sheet_name="GSWorkerWorkingHours", index=False)
            wo_data.to_excel(writer, sheet_name="WorkOrderData", index=False)

        # Planner Skills
        planner = pd.DataFrame([
            {"Worker ID": "W1", "Process": "Sanding", "Planner Verified Skill Level": 6.0},
            {"Worker ID": "W2", "Process": "Sanding", "Planner Verified Skill Level": 4.0},
        ])
        planner.to_excel(fixture_dir / "02_Planner_Skills" / "planner.xlsx", index=False)

        # Item Master
        item_m = pd.DataFrame([
            {"Item Number": "2SN001", "Item Number (Size Adjusted)": "2SN001", "Process": "Sanding", "Material": "Silver"},
        ])
        item_m.to_excel(fixture_dir / "09_Item_Master" / "item_master.xlsx", index=False)

        # QC Tickets with 42-column style names
        qc = pd.DataFrame([
            {"QualityOrderId": "Q1", "WO": "WO100", "ItemId": "2SN001", "RoundNo": 1, "QCStatus": "Pass", "QCQty": 20,
             "ExpectedInspectionQty": 20, "CreatedDateTime": "2026-08-02 10:00:00", "ValidatedDateTime": "1900-01-01 00:00:00"},
        ])
        qc.to_excel(fixture_dir / "04_QC_Tickets" / "qc.xlsx", index=False)

        result = run_pipeline(Config(input_dir=fixture_dir, output_dir=fixture_dir / "out", incomplete_month=""))
        data = result["data"]
        self.assertEqual(len(data), 3)
        # WO100 total hours = 4.0. W1 has 3.0/4.0 * 20 = 15.0 pieces. W2 has 1.0/4.0 * 20 = 5.0 pieces.
        w1_wo100 = data[(data.Reference == "WO100") & (data.Worker == "W1")].iloc[0]
        w2_wo100 = data[(data.Reference == "WO100") & (data.Worker == "W2")].iloc[0]
        self.assertAlmostEqual(w1_wo100["Qty Doing"], 15.0)
        self.assertAlmostEqual(w2_wo100["Qty Doing"], 5.0)
        self.assertAlmostEqual(w1_wo100["Total Actual Hours"], 3.0)
        self.assertAlmostEqual(w2_wo100["Total Actual Hours"], 1.0)
        # QC was reconstructed
        fpy = result["tables"]["FPY Semi"]
        self.assertFalse(fpy.empty)
        self.assertAlmostEqual(fpy["ObservedFPY"].iloc[0], 1.0)


if __name__ == "__main__":
    unittest.main()
