"""Constant-memory XLSX renderer for the full QC ticket and worker detail.

Artifact Tool cannot export this ~million-cell workbook within its Node heap,
so this renderer writes rows sequentially while preserving native Excel filters.
"""
from __future__ import annotations

import json
import math
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .semi_scope import EXCLUDED_STATUS

HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="17365D")
BODY_FONT = Font(name="Arial", size=10, color="1F2937")
NOTE_FONT = Font(name="Arial", size=10, italic=True, color="1F2937")
WARNING_FONT = Font(name="Arial", size=10, color="A61C00")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
BODY_ALIGN = Alignment(vertical="center")
SKU_HEADERS = [
    "SKU / Semi", "Process", "First Pass Yield (40%)",
    "Part & Mechanism (25%)", "Material Design Process 20%",
    "Stone 10%", "Learning 5%", "Total Score10", "Contribute/10",
    "% contribute", "First Pass Yield Summary", "FPY 40% confident", "% confident overall", "Semi scope status",
]
SKU_FACTORS = [
    ("QualityFactor_Approved", 0.40),
    ("PartMechanismFactor_Approved", 0.25),
    ("MaterialDesignFactor_Approved", 0.20),
    ("StoneFactor_Approved", 0.10),
    ("LearningFactor_Approved", 0.05),
]


def _number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _sku_score_rows(headers: list[str], records: list[list]) -> list[list]:
    ix = {name: pos for pos, name in enumerate(headers)}
    result = []
    for row in records:
        get = lambda name: row[ix[name]]
        optional = lambda name: row[ix[name]] if name in ix else None
        if optional("DifficultyStatus") == EXCLUDED_STATUS:
            result.append([get("SemiBOM"), get("Process"), None, None, None, None, None,
                           None, None, 0.0, "Excluded from Semi scoring by code rule", None, None, EXCLUDED_STATUS])
            continue
        approved = [get(name) if _number(get(name)) else None for name, _ in SKU_FACTORS]
        factors = approved.copy()
        if factors[0] is None and _number(get("QualityFactor_Diagnostic")):
            factors[0] = get("QualityFactor_Diagnostic")
        if factors[2] is None and _number(optional("MaterialDesignSourceScore")):
            factors[2] = optional("MaterialDesignSourceScore")
        if factors[3] is None and _number(get("StoneSourceScore")):
            factors[3] = get("StoneSourceScore")
        if any(score is not None and not 0 <= score <= 10 for score in factors):
            raise ValueError("Semi factor score outside 0–10")
        available = [(score, weight) for score, (_, weight) in zip(factors, SKU_FACTORS) if score is not None]
        contribution = sum(score * weight for score, weight in available) if available else None
        coverage = sum(weight for _, weight in available)
        published = get("FinalTechnicalComplexity")
        if _number(published) and any(score is None for score in approved):
            raise ValueError("Published final Semi score lacks all five approved factors")
        if _number(published) and not math.isclose(published, sum(score * weight for score, (_, weight) in zip(approved, SKU_FACTORS)), rel_tol=0, abs_tol=1e-7):
            raise ValueError("Published final Semi score does not reconcile to five factors")
        # Missing factors are omitted, then remaining weights are normalized
        # back to a 0–10 score. Contribution retains the unnormalized sum.
        score10 = contribution / coverage if coverage else None
        observed = get("ObservedFPY")
        if _number(observed):
            passed, inspected, n = get("PassQty1"), get("InspectedQty1"), get("QCEvidenceN")
            summary = f"{observed:.1%} observed"
            if _number(passed) and _number(inspected):
                summary += f" ({passed:g}/{inspected:g} pieces)"
            if _number(n):
                summary += f"; {n:g} QC WOs"
            summary += f"; {get('FirstPass_Calibration_Status') or 'NOT_VALIDATED'}"
        else:
            summary = "No eligible first-pass QC"
        model_confidences = []
        missing_model_confidence = False
        quality_confidence = None
        if factors[0] is not None:
            # Confidence must belong to the exact Quality score displayed.
            if approved[0] is not None:
                confidence = optional("QualityConfidencePct")
                status = optional("QualityConfidenceStatus")
                valid_status = {"ESTIMATED_VALIDATED"}
            else:
                confidence = get("RaschConfidencePct")
                status = get("RaschConfidenceStatus")
                valid_status = {"CONDITIONAL_DIAGNOSTIC", "ESTIMATED_VALIDATED"}
            if _number(confidence) and 0 <= confidence <= 1 and status in valid_status:
                model_confidences.append(confidence)
                quality_confidence = confidence
            else:
                missing_model_confidence = True
        for index, prefix in [(1, "PartMechanism"), (2, "MaterialDesign"), (3, "Stone"), (4, "Learning")]:
            if factors[index] is None or optional(prefix + "FactorMethod") != "MODEL_ESTIMATE":
                continue
            confidence = optional(prefix + "ConfidencePct")
            if _number(confidence) and 0 <= confidence <= 1 and optional(prefix + "ConfidenceStatus") == "ESTIMATED_VALIDATED":
                model_confidences.append(confidence)
            else:
                missing_model_confidence = True
        # Conservative completeness-adjusted index, not a calibrated joint
        # probability: the least precise visible model times weight coverage.
        confidence = None if missing_model_confidence else coverage * min(model_confidences, default=1.0)
        result.append([
            get("SemiBOM"), get("Process"), *factors, score10, contribution,
            coverage, summary, quality_confidence, confidence, optional("SemiScopeStatus") or "IN_SCOPE",
        ])
    return result


def _factor_provenance_rows(headers: list[str], records: list[list]) -> list[list]:
    """Expose which Semi factors are source/rubric values versus estimates."""
    ix = {name: pos for pos, name in enumerate(headers)}
    result = []
    for row in records:
        get = lambda name: row[ix[name]]
        optional = lambda name: row[ix[name]] if name in ix else None
        semi, process = get("SemiBOM"), get("Process")
        if optional("DifficultyStatus") == EXCLUDED_STATUS:
            for name, weight in [("Quality", .40), ("Part & Mechanism", .25),
                                 ("Material Design Process", .20), ("Stone", .10), ("Learning", .05)]:
                result.append([semi, process, name, weight, None, EXCLUDED_STATUS,
                               None, EXCLUDED_STATUS, "Excluded from Semi scoring by code rule"])
            continue
        quality = get("QualityFactor_Diagnostic")
        quality_estimated = _number(quality)
        quality_confidence = get("RaschConfidencePct")
        if not (quality_estimated and _number(quality_confidence)
                and get("RaschConfidenceStatus") in {"CONDITIONAL_DIAGNOSTIC", "ESTIMATED_VALIDATED"}):
            quality_confidence = None
        result.append([semi, process, "Quality", 0.40,
                       quality if quality_estimated else None,
                       "QC_RASCH_MODEL" if quality_estimated else "NOT_ESTIMATED",
                       quality_confidence,
                       get("RaschConfidenceStatus"),
                       "Within ±1 point of Quality estimate; conditional on selected QC models"])
        for name, weight, approved_column, pending_method in [
            ("Part & Mechanism", 0.25, "PartMechanismFactor_Approved", "ENGINEERING_RUBRIC_PENDING"),
            ("Learning", 0.05, "LearningFactor_Approved", "RAMP_UP_MODEL_PENDING"),
        ]:
            score = get(approved_column)
            result.append([semi, process, name, weight,
                           score if _number(score) else None,
                           "APPROVED_ENGINEERING_INPUT" if _number(score) else pending_method,
                           None, "NOT_APPLICABLE_FIXED_INPUT" if _number(score) else "NOT_ESTIMATED",
                           "If estimated by a model, provide factor-specific uncertainty before publication"])
        material = get("MaterialDesignFactor_Approved")
        source_material = optional("MaterialDesignSourceScore")
        approved_material = _number(material)
        material_score = material if approved_material else source_material if _number(source_material) else None
        result.append([semi, process, "Material Design Process", 0.20, material_score,
                       "APPROVED_ENGINEERING_INPUT" if approved_material else "SOURCE_RULE_CALCULATION_PROVISIONAL" if material_score is not None else "MISSING_SOURCE",
                       None, "NOT_APPLICABLE_FIXED_INPUT" if material_score is not None else "NOT_ESTIMATED",
                       "Source 0.4×Material + 0.4×Design + 0.2×Process; product type matched to Item Master"])
        stone = get("StoneFactor_Approved")
        source_stone = get("StoneSourceScore")
        approved_stone = _number(stone)
        score = stone if approved_stone else source_stone if _number(source_stone) else None
        result.append([semi, process, "Stone", 0.10, score,
                       "APPROVED_ENGINEERING_INPUT" if approved_stone else "SOURCE_RULE_CALCULATION_PROVISIONAL" if score is not None else "MISSING_SOURCE",
                       None, "NOT_APPLICABLE_FIXED_INPUT" if score is not None else "NOT_ESTIMATED",
                       "Source Stone score is calculated from specifications, not a fitted regression"])
    return result


def _cell(ws, value, *, header=False, title=False, italic=False, pct=False, warning=False):
    cell = WriteOnlyCell(ws, value=value)
    if header:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    elif title:
        cell.font = TITLE_FONT
    else:
        cell.font = WARNING_FONT if warning else NOTE_FONT if italic else BODY_FONT
        cell.alignment = BODY_ALIGN
        if pct and value is not None:
            cell.number_format = "0.0%"
        elif isinstance(value, float):
            cell.number_format = "0.000"
    return cell


def _sheet(wb, name: str, title: str, note: str, headers: list[str], records: list[list]) -> None:
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"
    ws.append([])
    ws.append([_cell(ws, title, title=True)])
    ws.append([_cell(ws, note, italic=True)])
    ws.append([])
    ws.append([_cell(ws, h, header=True) for h in headers])
    ws.row_dimensions[5].height = 27
    for row in records:
        ws.append([_cell(ws, value, pct=headers[i] in {"RoundFPY", "TicketSnapshotFPY", "ObservedFPY", "RaschConfidencePct", "Weight", "% contribute", "FPY 40% confident", "% confident overall", "% confident", "DifficultyCoverage", "EstimatePrecisionWithin20Pct", "CalibrationWithin20Pct", "WAPE", "Within20Pct", "PredictionInterval90Coverage"}, warning=isinstance(value, str) and value in {"BLOCKED", "NOT_VALIDATED", "NOT_ESTIMATED", "ITEM_MISMATCH"}) for i, value in enumerate(row)])
    ws.auto_filter.ref = f"A5:{get_column_letter(len(headers))}{len(records)+5}"
    for col, h in enumerate(headers, start=1):
        width = 17
        if any(token in h for token in ("SemiBOM", "SemiItem", "QualityOrderId", "SizeAdjustedGroup")):
            width = 25
        if any(token in h for token in ("Name", "Status", "SourceFile")):
            width = 29
        if any(token in h for token in ("Score", "Difficulty", "FPY", "Effect")):
            width = 20
        ws.column_dimensions[get_column_letter(col)].width = width
    print(f"{name}: {len(records)} rows", flush=True)


def build(destination: Path) -> Path:
    manifest = json.loads((destination / "manifest.json").read_text(encoding="utf-8"))
    wb = Workbook(write_only=True)
    semi_headers = manifest["sheet_columns"]["Semi difficulty"]
    semi_records = json.loads((destination / "Semi difficulty.json").read_text(encoding="utf-8"))
    sku_records = _sku_score_rows(semi_headers, semi_records)
    summary = [
        ["Metric", "Value"],
        ["SemiBOMs", manifest["semi_boms"]],
        ["Excluded SemiBOMs by code rule", manifest["excluded_semi_scope"]],
        ["Approved final Semi difficulties", manifest["approved_final_difficulty_rows"]],
        ["SemiBOMs with Stone source score", manifest["semi_with_stone_source_score"]],
        ["SemiBOMs with Material Design Process source score", manifest["semi_with_material_design_source_score"]],
        ["SemiBOMs without Stone source score", manifest["semi_without_stone_source_score"]],
        ["SemiBOMs with modeled Quality estimate", manifest["semi_with_diagnostic_quality_estimate"]],
        ["SemiBOMs with quoted minutes", manifest["semi_with_quoted_minutes"]],
        ["Stone FG-context source rows", manifest["stone_source_rows"]],
        ["QC tickets", manifest["qc_tickets"]],
        ["QC tickets linked to core round", manifest["qc_tickets_linked_to_core_rounds"]],
        ["QC tickets with first-round FPY", manifest["qc_tickets_with_round_fpy"]],
        ["QC tickets with ticket-snapshot FPY", manifest["qc_tickets_with_snapshot_fpy"]],
        ["Ticket / production worker rows", manifest["ticket_worker_rows"]],
        ["Diagnostic Rasch worker effects", manifest["rasch_worker_effects"]],
        ["Per-Semi Rasch confidence percentages", manifest["rasch_confidence_pct_rows"]],
        ["SKU rows with available-factor Score10", sum(row[7] is not None for row in sku_records)],
        ["SKU scored rows with FPY 40% confidence", sum(row[7] is not None and row[11] is not None for row in sku_records)],
        ["SKU scored rows with overall confidence index", sum(row[7] is not None and row[12] is not None for row in sku_records)],
        ["SKU rows with all five factors", sum(math.isclose(row[9], 1.0) for row in sku_records)],
    ]
    _sheet(wb, "Summary", "Semi difficulty, Stone and QC evidence", "SKU Score10 uses available factors; confidence is reduced by missing weight. Approved final score remains separately gated.", summary[0], summary[1:])
    _sheet(
        wb, "SKU Semi Score", "SKU / Semi score (0–10)",
        "Excluded codes have no score. In-scope missing factors are skipped; FPY confidence and completeness-adjusted overall confidence are separate.",
        SKU_HEADERS, sku_records,
    )
    _sheet(
        wb, "Factor provenance", "Semi factor methods and confidence",
        "Confidence applies to fitted estimates. Fixed source/rubric inputs have no statistical confidence percentage; missing models stay blank.",
        ["SKU / Semi", "Process", "Factor", "Weight", "Score10", "Method", "% confident", "Confidence status", "Confidence basis"],
        _factor_provenance_rows(semi_headers, semi_records),
    )
    notes = {
        "Semi difficulty": "One row per SemiBOM. Stone source and Rasch group difficulty are diagnostic inputs; blank final score is not zero.",
        "Material Design source": "Original Material/Design/Process engineering rows. SelectionStatus identifies the product-type-matched row used per SemiBOM.",
        "Stone source": "New Stone workbook, one row per FG–Semi source context. Filter SemiBOM to inspect provenance.",
        "QC tickets": "One row per QualityOrderId. Ticket snapshot FPY and matched first-round FPY have different grains.",
        "Ticket workers": "One row per ticket and attributed production worker; team QC results are shared, not allocated.",
        "Rasch workers": "Process/branch fitted worker effects are diagnostic; calibration is not validated.",
        "Quote source": "Original BOM quote rows. Conflicts and nonpositive minutes are retained but not used as calibrated inputs.",
    }
    for name, headers in manifest["sheet_columns"].items():
        records = semi_records if name == "Semi difficulty" else json.loads((destination / f"{name}.json").read_text(encoding="utf-8"))
        _sheet(wb, name, name, notes[name], headers, records)
        del records
    definitions = [
        ["Field", "Definition / interpretation"],
        ["FinalTechnicalComplexity", "Approved 0–10 Semi difficulty. Blank when Rasch and factor gates are not met; blank is not zero."],
        ["SKU sheet First Pass Yield (40%)", "QC-model estimate: 80% first-pass difficulty plus 20% rework difficulty inside Quality. Selected penalized models adjust for worker effects where supported. Provisional, not raw FPY."],
        ["SKU sheet Material Design Process 20%", "Source workbook composite: 40% Material Score + 40% Design Score + 20% Process Score, all 0–10. Product type must match Item Master. Provisional fixed engineering input, not a fitted regression or approved final factor."],
        ["SKU sheet Stone 10%", "New Stone workbook's Final Semi Stone Score when an approved Stone factor is not yet available. Provisional source score."],
        ["SKU sheet Total Score10", "Available-factor score on 0–10: sum(score × weight) / sum(available weights). Blank if no factor is available. Partial or provisional when coverage is below 100% or factor validation is pending."],
        ["SKU sheet Contribute/10", "Unnormalized weighted sum of visible available factor estimates, including provisional Quality and Stone."],
        ["SKU sheet % contribute", "Sum of weights for visible available factor estimates, from 0% to 100%. This is completeness, not statistical confidence or approval."],
        ["SKU sheet First Pass Yield Summary", "Observed first-round QC FPY, pass/inspected pieces and WO evidence count. It is diagnostic and does not bypass Quality approval."],
        ["SKU sheet FPY 40% confident", "Conditional QC/Rasch precision belonging only to the displayed 40% Quality estimate. Blank if that factor or its matching precision is unavailable; does not inherit fixed-factor coverage."],
        ["SKU sheet % confident overall", "Completeness-adjusted confidence index = available weight × lowest confidence of the model-estimated factors used. Fixed-only rows start at 100% before missing-weight reduction. Blank if any included model estimate lacks its own confidence. This is a heuristic index, not a calibrated probability that total Score10 is correct."],
        ["Semi scope status", "EXCLUDED_SEMI_SCOPE means the item starts with an excluded prefix or its Semi/BOM code ends -01 or -02. Such rows remain for reconciliation but have no factor, total score or confidence."],
        ["SKU confidence example", "Quality model confidence 60% plus Stone score gives 50% available weight, so FPY 40% confident stays 60% and overall confidence index is 60% × 50% = 30%."],
        ["Factor provenance", "One row per Semi and factor. Method distinguishes fitted estimates from source/rubric scores and missing factors. Confidence percent is populated only for a fitted factor with an explicit uncertainty calculation."],
        ["Fixed versus estimated", "Quoted minutes, BOM/specification fields and approved rubric scores are fixed inputs for this run. QC Quality and worker/time effects are fitted estimates; their precision must be quantified independently."],
        ["Learning model gate", "The 5% ramp-up factor has no fitted model in the current run. If estimated from time/effort trajectories, output a separate confidence percentage and interval before using it in the final Semi score."],
        ["RaschInterval90Lower/Upper", "Conditional 90% interval for combined Quality estimate from the same local-information approximation. Excludes selection, worker-effect and model-specification uncertainty."],
        ["QualityFactor_Approved", "40% of final score: calibrated first-pass/FPY difficulty plus approved rework trajectory (80/20 inside Quality under the current specification). Both calibration gates must pass."],
        ["PartMechanismFactor_Approved", "25% of final score; rubric subitems P1, P2, P3. Direct SM_Casting and SM_ACJ PCS part count is evidence, not the approved score."],
        ["MaterialDesignFactor_Approved", "20% of final score; rubric subitems M1, M2, M3. Material group count is evidence, not the approved score."],
        ["MaterialDesignSourceScore", "Material Design Process source composite from the new workbook, shown provisionally in SKU Score10. Blank when source or product-type match is missing."],
        ["StoneFactor_Approved", "10% of final score; rubric subitems S1, S2, S3. Requires approved rubric, scorer, approver and version."],
        ["LearningFactor_Approved", "5% of final score; rubric subitems L1, L2, L3. Requires approved engineering assessment."],
        ["StoneSourceScore", "Final Semi Stone Score from new Stone workbook. Unapproved source measure, not released Stone factor or total difficulty."],
        ["QuotedMinutesPerSemi", "User-provided planned minutes for one Semi. A QuotedHours input is converted by multiplying by 60. This is Work Content input, not a difficulty score."],
        ["QuoteBasis", "PER_SEMI_UNIT. Quoted time is for one Semi, not the total duration of a WO or batch."],
        ["Worker/time separation", "Worker Effect, recorded time and residual Error belong to the time model. They are not components of the five-factor Semi difficulty score."],
        ["Time model unit rule", "Convert 0–10 difficulty to a calibrated minute adjustment before combining with quoted minutes; do not add score points directly to minutes or count difficulty twice."],
        ["RaschGroupDifficulty_Diagnostic", "First-pass Rasch model-group difficulty (0–10), not validated and not approved for decisions."],
        ["RaschConfidencePct", "Conditional QC-model precision for the 40% Quality factor only. Estimated from local first-pass and rework round information; excludes model-selection, worker-effect and specification uncertainty. Blank without a modeled Quality factor."],
        ["TicketSnapshotFPY", "FirstPassQty / (FirstPassQty + FirstFailQty) if denominator > 0. A ticket can be a cumulative snapshot; do not sum ticket quantities."],
        ["RoundFPY", "Matched first WO round PassQty / InspectedQty. Repeated for tickets in the same round; not independent observations."],
        ["Rasch_Worker_Effect", "Diagnostic process/branch worker logit effect, not a ticket score or approved worker skill score."],
        ["WorkerAttributionStatus", "Production worker attribution from core WO-round join. Team results are not split into individual successes."],
        ["QCWorkerWrkCtrId", "QC inspector from source ticket, distinct from attributed production Worker."],
        ["Stone provenance", "input_data/06_Engineering_Factors/Stone/Semi Stone Score.xlsx; join on BOM item number."],
        ["Material Design provenance", "input_data/06_Engineering_Factors/Material Design Process/Material Design Process Score.xlsx; join on Semi BOM Item, Semi Item and Process; select duplicate context by Item Master Product Type."],
        ["QC provenance", "input_data/04_QC_Tickets/QC data.xlsx; join to completed core run on WO + RoundNo."],
        ["Core provenance", "Do kho SKU.csv, QC rounds.csv, QC round workers.csv, Rasch first pass workers.csv, Rasch rework workers.csv."],
        ["Quoted time input", "Current input_data/08_Quoted Hours/Quoted Hours.xlsx has Minutes per Semi, keyed by BOM number. Conflicting and nonpositive quotes remain auditable and are not used for calibration."],
    ]
    _sheet(wb, "Definitions", "Definitions and sources", "Scores are current source diagnostics unless explicitly marked approved.", definitions[0], definitions[1:])
    target = destination / "semi_rasch_stone_detail.xlsx"
    wb.save(target)
    return target


if __name__ == "__main__":
    print(build(Path(sys.argv[1]).resolve()), flush=True)
