"""Create one blank, validated workbook in each input-data folder."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from .create_template import SCHEMAS
from .workbook_io import style_table

FOLDERS = {
    "Production": "01_Production",
    "Planner Skills": "02_Planner_Skills",
    "Item Mapping": "03_Item_Mapping",
    "QC Tickets": "04_QC_Tickets",
    "Touch Events": "05_Touch_Events",
    "Engineering Factors": "06_Engineering_Factors",
    "Pilot Log": "07_Pilot_Log",
    "Reference Master": "08_Reference_Master",
    "Item Master": "09_Item_Master",
    "Worker Master": "10_Worker_Master",
}


def create_folder_templates(root=Path("input_data"), overwrite=False):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    created = []
    for dataset, fields in SCHEMAS.items():
        folder = root / FOLDERS[dataset]
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{dataset.replace(' ', '_')}_template.xlsx"
        if path.exists() and not overwrite:
            continue
        wb = Workbook()
        ws = wb.active
        ws.title = dataset[:31]
        headers = [header for header, _ in fields]
        ws.append(headers)
        style_table(ws, 1, len(headers), list(SCHEMAS).index(dataset) + 1)
        for column, (header, definition) in enumerate(fields, 1):
            letter = get_column_letter(column)
            ws.cell(1, column).comment = Comment(definition, "Skill ID")
            ws.column_dimensions[letter].width = min(max(len(header) + 3, 21), 35)
            is_date = header in ["RAF Month", "QC_Start", "QC_Stop", "Worker_Start", "Worker_Stop", "EffectiveFrom", "AssignmentTime"]
            is_number = header in ["Qty Doing", "Total Actual Hours", "QCQty", "ExpectedQty", "RoundNo", "Score", "ConfidencePct", "Planner Verified Skill Level", "HandedFailQty", "RecoveredQty"]
            for row in range(2, 102):
                cell = ws.cell(row, column)
                cell.font = Font(name="Arial", size=11)
                cell.number_format = "yyyy-mm-dd hh:mm:ss" if is_date else ("0.0%" if header == "ConfidencePct" else "0.000" if is_number else "@")
            options = None
            if header == "QCStatus": options = '"Pass,Fail"'
            elif header == "Factor": options = '"Quality,PartMechanism,MaterialDesignProcess,Stone,Learning"'
            elif header in ["Approved", "FollowedTop3"]: options = '"TRUE,FALSE"'
            elif header == "Arm": options = '"Control,Intervention"'
            elif header == "DataQualityStatus": options = '"Valid,Pending,Exception"'
            elif header == "Method": options = '"FIXED_INPUT,MODEL_ESTIMATE"'
            elif header == "ConfidenceStatus": options = '"ESTIMATED_VALIDATED,CONDITIONAL_DIAGNOSTIC,NOT_ESTIMATED,NOT_APPLICABLE"'
            validation = None
            if options:
                validation = DataValidation(type="list", formula1=options, allow_blank=True)
            elif header in ["Score", "Planner Verified Skill Level"]:
                validation = DataValidation(type="decimal", operator="between", formula1=0, formula2=10, allow_blank=True)
            elif header == "ConfidencePct":
                validation = DataValidation(type="decimal", operator="between", formula1=0, formula2=1, allow_blank=True)
            elif header in ["RoundNo", "QCQty", "ExpectedQty"]:
                validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1=0 if header == "QCQty" else 1, allow_blank=True)
            elif is_number:
                validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True)
            if validation:
                validation.showErrorMessage = True
                validation.errorTitle = "Invalid entry"
                validation.error = "Check the column comment and input_data/README.md."
                ws.add_data_validation(validation)
                validation.add(f"{letter}2:{letter}1048576")
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}101"
        ws.freeze_panes = "C2"
        ws.sheet_view.showGridLines = False
        wb.save(path)
        created.append(path)
    return created


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path("input_data"))
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    for path in create_folder_templates(args.input_dir, args.overwrite):
        print(f"Created {path.resolve()}")
