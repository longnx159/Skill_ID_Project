"""Durable run lifecycle and English reports; no model approval is inferred."""
from __future__ import annotations

from dataclasses import asdict, replace
from datetime import datetime, timezone
import hashlib
import json
import logging
import os
from pathlib import Path
import platform
import shutil
import socket
import time
import uuid

import numpy as np
import pandas as pd


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_json(value):
    if isinstance(value, dict):
        return {str(k): clean_json(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean_json(v) for v in value]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.generic):
        return clean_json(value.item())
    if value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, float) and not np.isfinite(value):
        return None
    return value


def atomic_json(path, value):
    path = Path(path)
    temp = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    try:
        temp.write_text(json.dumps(clean_json(value), indent=2, ensure_ascii=False,
                                   allow_nan=False, default=str), encoding="utf-8")
        # Windows indexers/antivirus can briefly hold the destination open.
        # Retry only a bounded sharing/permission failure; never overwrite via
        # a non-atomic fallback or hide a persistent permission problem.
        for attempt in range(5):
            try:
                temp.replace(path)
                break
            except PermissionError:
                if attempt == 4:
                    raise
                time.sleep(.05 * (2 ** attempt))
    finally:
        temp.unlink(missing_ok=True)


def records(tables, name, required):
    frame = tables.get(name, pd.DataFrame())
    return frame if set(required).issubset(frame.columns) else pd.DataFrame(columns=required)


class WarningCollector(logging.Handler):
    def __init__(self, messages):
        super().__init__(logging.WARNING)
        self.messages = messages

    def emit(self, record):
        message = record.getMessage()
        if message not in self.messages:
            self.messages.append(message)


class RunSession:
    """A unique child of the requested output root; never overwrites prior runs."""

    def __init__(self, config):
        self.config = config
        self.root = Path(config.output_dir).resolve()
        self.run_id = datetime.now(timezone.utc).strftime("run_%Y%m%dT%H%M%S_%fZ_") + uuid.uuid4().hex[:8]
        self.path = self.root / self.run_id
        self.path.mkdir(parents=True, exist_ok=False)
        self.staging = self.path / "artifacts.partial"
        self.staging.mkdir()
        self.started = time.perf_counter()
        self.active = None
        self.logger = logging.getLogger("pipeline")
        self.old_level = self.logger.level
        self.logger.setLevel(logging.INFO)
        self.log_handler = logging.FileHandler(self.path / "run.log", encoding="utf-8")
        self.log_handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s"))
        self.summary = {
            "schema_version": 1, "run_id": self.run_id, "output_directory": str(self.path),
            "execution_status": "RUNNING", "decision_status": "DIAGNOSTIC_ONLY",
            "started_at": utc_now(), "finished_at": None, "duration_seconds": None,
            "owner_pid": os.getpid(), "owner_host": socket.gethostname(),
            "timing_scope": "Start through data export; report finalization is excluded",
            "configuration": clean_json(asdict(config)), "stages": [], "warnings": [],
            "data": {}, "qc": {}, "model_health": [], "validation": [], "gates": [],
            "sources": [], "artifacts": [], "exit_code": None,
            "comparison": {"status": "NO_BASELINE"},
            "planner_policy": "2026-08-03 verified baseline applied retrospectively; not historical as-of skill evidence",
            "metric_definitions": {"WAPE minutes": "Ratio: sum absolute minute errors / sum actual minutes",
                                   "RMSE": "Log-time units", "MAE": "Log-time units",
                                   "FPY": "Total first-round passed pieces / total first-round inspected pieces"},
            "acceptance_thresholds": "NOT_DEFINED",
        }
        self.collector = WarningCollector(self.summary["warnings"])
        self.logger.addHandler(self.log_handler)
        self.logger.addHandler(self.collector)
        effective = clean_json(asdict(config))
        for key in ("input_dir", "output_dir", "raw_path", "planner_path"):
            effective.pop(key, None)
        self.summary["configuration_sha256"] = hashlib.sha256(json.dumps(effective, sort_keys=True).encode()).hexdigest()
        self.summary["code_sha256"] = {p.name: file_hash(p) for p in Path(__file__).parent.glob("*.py")}
        import openpyxl
        self.summary["runtime"] = {"python": platform.python_version(), "pandas": pd.__version__,
                                   "numpy": np.__version__, "openpyxl": openpyxl.__version__}
        try:
            self.checkpoint()
        except BaseException:
            self.close()
            raise

    def checkpoint(self):
        atomic_json(self.path / "run_summary.json", self.summary)

    def start_stage(self, name):
        self.finish_stage()
        stage = {"stage": name, "status": "RUNNING", "started_at": utc_now(), "duration_seconds": None}
        self.summary["stages"].append(stage)
        self.active = (stage, time.perf_counter())
        self.logger.info("Starting %s", name)
        self.checkpoint()

    def finish_stage(self, status="SUCCEEDED"):
        if self.active:
            stage, start = self.active
            stage.update(status=status, duration_seconds=round(time.perf_counter() - start, 3))
            self.active = None
            self.checkpoint()

    def call(self, name, function, *args, **kwargs):
        self.start_stage(name)
        result = function(*args, **kwargs)
        self.finish_stage()
        return result

    def snapshot_inputs(self, template=None):
        from .pipeline_v053 import discover_input_files, FOLDER_DATASETS
        if template:
            files = {"Template": [Path(template)]}
        else:
            if not Path(self.config.input_dir).is_dir():
                raise FileNotFoundError(f"Input folder does not exist: {self.config.input_dir}")
            files = discover_input_files(self.config.input_dir)
        snapshot = self.path / "inputs"
        snapshot.mkdir()
        for dataset, paths in files.items():
            folder = snapshot / (FOLDER_DATASETS[dataset][0] if dataset != "Template" else "template")
            for source in paths:
                if source.suffix.lower() == ".xls":
                    try:
                        import xlrd  # noqa: F401
                    except ImportError as exc:
                        raise ValueError(f"Legacy .xls requires xlrd; convert to .xlsx: {source}") from exc
                folder.mkdir(exist_ok=True)
                target = folder / source.name
                before = file_hash(source)
                shutil.copyfile(source, target)
                if before != file_hash(target) or before != file_hash(source):
                    raise ValueError(f"Input changed while snapshotting: {source}")
                self.summary["sources"].append({"path": str(source.resolve()), "sha256": before,
                                                "snapshot": str(target.relative_to(self.path)),
                                                "size_bytes": target.stat().st_size, "dataset": dataset})
        self.checkpoint()
        snapshot_template = snapshot / "template" / Path(template).name if template else None
        return replace(self.config, input_dir=snapshot), snapshot_template

    def collect(self, result):
        tables, data = result["tables"], result["data"]
        quality = records(tables, "Data Quality", ["Check", "N"])
        checks = dict(zip(quality.Check, quality.N))
        raw = int(checks.get("Production input rows", len(data)))
        basic = int(checks.get("Production retained after basic filters", len(data)))
        self.summary["data"] = {
            "input_production_rows": raw, "removed_by_basic_filters_and_cutoff": raw - basic,
            "unique_quarantined_rows": basic - len(data), "retained_production_rows": len(data),
            "reconciled": raw == (raw - basic) + (basic - len(data)) + len(data),
            "grain": "Reference + Worker + Item Number + RAF Month + RoundNo",
            "reason_counts_may_overlap": quality.to_dict("records"),
            "estimated_quantity_rows": int(data.get("Qty Doing Source", pd.Series("", index=data.index)).str.contains("estimated", na=False).sum()),
            "unknown_round_rows": int(data.RoundNo.isna().sum()),
        }
        first = records(tables, "WO vong dau", ["WO", "PassQty", "InspectedQty"])
        rounds = records(tables, "QC rounds", ["WO", "RoundNo"])
        links = records(tables, "QC round workers", ["WO", "RoundNo", "WorkerAttributionStatus"])
        denominator = first.InspectedQty.sum()
        self.summary["qc"] = {
            "availability": "AVAILABLE" if "WO" in tables.get("QC rounds", pd.DataFrame()).columns else "NOT_AVAILABLE",
            "eligible_rounds": len(rounds.drop_duplicates(["WO", "RoundNo"])),
            "first_round_wos": first.WO.nunique(),
            "later_rounds": int(rounds.RoundNo.ge(2).sum()),
            "single_worker_rounds": len(links.loc[links.WorkerAttributionStatus.eq("Single worker in WO-round")].drop_duplicates(["WO", "RoundNo"])),
            "observed_fpy": float(first.PassQty.sum() / denominator) if denominator else None,
            "scrap_status": "Not certified; missing quantities remain unknown",
        }
        if self.summary["qc"]["availability"] == "NOT_AVAILABLE":
            for key in ("eligible_rounds", "first_round_wos", "later_rounds", "single_worker_rounds"):
                self.summary["qc"][key] = None
        self.summary["qc"]["single_worker_attribution_pct"] = (
            100 * self.summary["qc"]["single_worker_rounds"] / self.summary["qc"]["eligible_rounds"]
            if self.summary["qc"]["eligible_rounds"] else None)
        trajectories = records(tables, "QC hanh trinh", ["WO"])
        self.summary["qc"]["trajectory_status_counts"] = (
            trajectories.RecoveryStatus.value_counts().to_dict()
            if "RecoveryStatus" in trajectories else "NOT_AVAILABLE")
        exceptions = records(tables, "QC can kiem", ["Reason"])
        self.summary["qc"]["exception_reason_counts"] = exceptions.Reason.value_counts().to_dict()
        # Keep the summary compact: full group-level network data stays in its CSV.
        network = tables["Connectivity"]
        self.summary["network"] = (network.groupby(["Process", "ConnectivityStatus"]).size().rename("group_count").reset_index().to_dict("records")
                                   if {"Process", "ConnectivityStatus"}.issubset(network.columns) else [])
        self.summary["cutoff"] = result["manifest"]["cutoff"]
        self.summary["validation"] = tables["Model Validation"].to_dict("records")
        self.summary["gates"] = tables["Go Live Gates"].to_dict("records")
        health = []
        for process, model in result["models"].items():
            health.append({"process": process, "branch": "Time", "fit_status": "CONVERGED" if model["converged"] else "NONCONVERGED",
                           "evidence_n": int(data.Process.eq(process).sum()), "iterations": model.get("iterations"),
                           "log_likelihood": model.get("log_likelihood"), "calibration_status": "NOT_VALIDATED",
                           "decision_eligibility": "BLOCKED"})
        processes = sorted(data.Process.dropna().unique())
        for name, branch in (("Rasch first pass groups", "First pass"), ("Rasch rework groups", "Rework")):
            frame = records(tables, name, ["Process", "Rasch_Converged", "Rasch_Evidence_N"])
            for process in processes:
                part = frame.loc[frame.Process.eq(process)]
                converged = not part.empty and bool(part.Rasch_Converged.all())
                health.append({"process": process, "branch": branch,
                               "fit_status": ("CONVERGED" if converged else "NONCONVERGED") if len(part) else "INSUFFICIENT_EVIDENCE",
                               "evidence_n": int(part.Rasch_Evidence_N.sum()),
                               "iterations": part.iloc[0].get("Rasch_Iterations") if len(part) else None,
                               "final_delta": part.iloc[0].get("Rasch_Final_Delta") if len(part) else None,
                               "calibration_status": "NOT_VALIDATED", "decision_eligibility": "BLOCKED"})
        self.summary["model_health"] = clean_json(health)
        for row in health:
            if row["fit_status"] == "NONCONVERGED":
                self.logger.warning("%s / %s is nonconverged; decision use is blocked", row["process"], row["branch"])
        if basic != len(data):
            self.logger.warning("%d production rows quarantined; inspect Production exceptions", basic - len(data))
        self.summary["actions"] = [
            "Review nonconverged or insufficient-evidence model branches before interpreting estimates.",
            "Resolve source exceptions and confirm mapping/round attribution.",
            "Complete calibration and Go Live Gates before production use.",
        ]
        self.compare()

    def compare(self):
        pointer = self.root / "latest_successful.json"
        if not pointer.exists():
            return
        try:
            previous_id = json.loads(pointer.read_text(encoding="utf-8"))["run_id"]
            if Path(previous_id).name != previous_id:
                raise ValueError("Invalid previous run ID")
            previous_path = self.root / previous_id
            manifest = json.loads((previous_path / "run_manifest.json").read_text(encoding="utf-8"))
            if manifest.get("execution_status") not in {"SUCCEEDED", "SUCCEEDED_WITH_WARNINGS"}:
                raise ValueError("Previous run is not complete")
            previous = json.loads((previous_path / "run_summary.json").read_text(encoding="utf-8"))
            summary_entry = next(a for a in manifest["artifacts"] if a["path"] == "run_summary.json")
            if file_hash(previous_path / "run_summary.json") != summary_entry["sha256"]:
                raise ValueError("Previous summary integrity check failed")
            reasons = []
            for key in ("configuration_sha256", "code_sha256", "runtime"):
                if previous.get(key) != self.summary.get(key):
                    reasons.append(key + " changed")
            old_sources = sorted((s["dataset"], s["sha256"]) for s in previous["sources"])
            new_sources = sorted((s["dataset"], s["sha256"]) for s in self.summary["sources"])
            if old_sources != new_sources:
                reasons.append("Source snapshot/cohort changed")
            comparison = {"previous_run_id": previous_id,
                          "status": "NOT_DIRECTLY_COMPARABLE" if reasons else "COMPARABLE",
                          "reasons": reasons}
            if not reasons:
                comparison["retained_row_change"] = self.summary["data"]["retained_production_rows"] - previous["data"]["retained_production_rows"]
                old_fpy, new_fpy = previous["qc"].get("observed_fpy"), self.summary["qc"]["observed_fpy"]
                comparison["fpy_change_percentage_points"] = (new_fpy - old_fpy) * 100 if old_fpy is not None and new_fpy is not None else None
                comparison["new_warnings"] = sorted(set(self.summary["warnings"]) - set(previous["warnings"]))
                comparison["resolved_warnings"] = sorted(set(previous["warnings"]) - set(self.summary["warnings"]))
                old_metrics = {row["Process"]: row for row in previous.get("validation", [])}
                comparison["wape_ratio_changes"] = {
                    row["Process"]: row["WAPE minutes"] - old_metrics[row["Process"]]["WAPE minutes"]
                    for row in self.summary["validation"] if row.get("WAPE minutes") is not None
                    and row["Process"] in old_metrics and old_metrics[row["Process"]].get("WAPE minutes") is not None}
            self.summary["comparison"] = comparison
        except (OSError, ValueError, KeyError, StopIteration, TypeError) as exc:
            self.summary["comparison"] = {"status": "BASELINE_UNAVAILABLE", "reason": str(exc)}
            self.logger.warning("Previous run comparison unavailable: %s", exc)

    def report_rows(self):
        rows = []
        for key, value in clean_json(self.summary).items():
            values = value if isinstance(value, list) else [value]
            for entry in values:
                text = json.dumps(entry, ensure_ascii=False, default=str) if isinstance(entry, dict) else str(entry if entry is not None else "NOT_AVAILABLE")
                rows.append({"Section": key.replace("_", " ").title(), "Details": text})
        return pd.DataFrame(rows)

    def render(self):
        s = clean_json(self.summary)
        lines = ["# Skill ID Run Summary", "", f"**Run ID:** {s['run_id']}",
                 f"**Execution:** {s['execution_status']}  ", f"**Decision status:** {s['decision_status']}",
                 "", "Results are diagnostic. Execution success does not authorize worker certification or automated assignment.", ""]
        for key, value in s.items():
            if key in {"run_id", "execution_status", "decision_status"}:
                continue
            lines.extend(["## " + key.replace("_", " ").title(), ""])
            if key == "model_health" and value:
                lines.extend(["| Process | Branch | Fit status | Evidence N | Iterations | Final delta | Calibration | Decision use |",
                              "|---|---|---|---:|---:|---|---|---|"])
                for row in value:
                    lines.append(f"| {row['process']} | {row['branch']} | {row['fit_status']} | {row['evidence_n']} | {row.get('iterations') or 'NOT_AVAILABLE'} | {row.get('final_delta', 'NOT_AVAILABLE')} | {row['calibration_status']} | {row['decision_eligibility']} |")
            elif key == "stages":
                lines.extend(["| Stage | Status | Seconds | Reason |", "|---|---|---:|---|"])
                for row in value:
                    lines.append(f"| {row['stage']} | {row['status']} | {row.get('duration_seconds', '—')} | {row.get('reason', '')} |")
            elif key == "validation" and value:
                lines.extend(["| Process | Status | Train N | Test N | WAPE | RMSE (log time) | MAE (log time) |", "|---|---|---:|---:|---:|---:|---:|"])
                for row in value:
                    wape = row.get("WAPE minutes")
                    wape = f"{wape:.2%}" if wape is not None else "NOT_AVAILABLE"
                    lines.append(f"| {row['Process']} | {row['Status']} | {row['Training N']} | {row['Test N']} | {wape} | {row.get('RMSE', 'NOT_AVAILABLE')} | {row.get('MAE', 'NOT_AVAILABLE')} |")
                lines.append("\nTrain/test windows, embargo and unseen-entity counts are retained in run_summary.json and Model Validation.csv.")
            elif key in {"warnings", "actions", "artifacts"}:
                lines.extend([f"- {entry}" for entry in value] or ["None."])
            elif key == "gates" and value:
                lines.extend(["| Gate | Status | Requirement |", "|---|---|---|"])
                lines.extend(f"| {r['Gate']} | {r['Status']} | {r['Requirement']} |" for r in value)
            elif key in {"data", "qc", "comparison"}:
                for name, entry in value.items():
                    if isinstance(entry, (dict, list)):
                        entry = json.dumps(entry, ensure_ascii=False)
                    lines.append(f"- **{name.replace('_', ' ').title()}:** {entry if entry is not None else 'NOT_AVAILABLE'}")
            elif isinstance(value, (dict, list)):
                lines.extend(["```json", json.dumps(value, indent=2, ensure_ascii=False, default=str), "```"])
            else:
                lines.append(str(value if value is not None else "NOT_AVAILABLE"))
            lines.append("")
        return "\n".join(lines)

    def finish(self, result=None, error=None):
        failed = error is not None
        status = "INTERRUPTED" if isinstance(error, KeyboardInterrupt) else "FAILED"
        if failed:
            self.summary["failed_stage"] = self.active[0]["stage"] if self.active else "Report finalization"
            self.finish_stage(status)
            self.summary["error"] = {"type": type(error).__name__, "message": str(error)}
            self.summary["decision_status"] = "BLOCKED"
        else:
            self.finish_stage()
            status = "SUCCEEDED_WITH_WARNINGS" if self.summary["warnings"] else "SUCCEEDED"
        self.summary.update(execution_status=status, finished_at=utc_now(),
                            duration_seconds=round(time.perf_counter() - self.started, 3),
                            exit_code=(130 if status == "INTERRUPTED" else 1) if failed else 0)
        if not failed:
            final = self.path / "artifacts"
            self.staging.rename(final)
            self.summary["artifacts"] = [str(p.relative_to(self.path)) for p in sorted(final.iterdir())]
            # Add the finalized summary to the existing workbook; same summary object as JSON/Markdown.
            from openpyxl import load_workbook
            from .workbook_io import style_table
            workbook_path = final / "Skill_ID_Ket_qua_chay_thu.xlsx"
            workbook = load_workbook(workbook_path)
            try:
                for name, frame in result["tables"].items():
                    expected_rows = len(frame) if len(frame.columns) else 1
                    if workbook[name[:31]].max_row != expected_rows + 1:
                        raise ValueError(f"Workbook row-count validation failed: {name}")
                sheet = workbook.create_sheet("Run Summary", 0)
                sheet.append(["Section", "Details"])
                for row in self.report_rows().itertuples(index=False, name=None):
                    sheet.append(row)
                style_table(sheet, sheet.max_row, 2, number=9999)
                sheet.column_dimensions["B"].width = 110
                workbook.save(workbook_path)
            finally:
                workbook.close()
        else:
            self.summary["artifacts"] = [str(p.relative_to(self.path)) for p in self.path.rglob("*") if p.is_file() and p.parent.name in {"artifacts", "artifacts.partial"}]
        self.checkpoint()
        (self.path / "run_summary.md").write_text(self.render(), encoding="utf-8")
        self.logger.info("Run %s: %s; report: %s", self.run_id, status, self.path / "run_summary.md")
        self.log_handler.flush()
        artifact_paths = [self.path / name for name in self.summary["artifacts"]]
        artifact_paths += [self.path / "run_summary.json", self.path / "run_summary.md", self.path / "run.log"]
        manifest = dict(result["manifest"]) if result else {}
        manifest.update({k: self.summary[k] for k in ("run_id", "execution_status", "decision_status", "configuration", "configuration_sha256", "code_sha256", "runtime", "sources", "started_at", "finished_at", "exit_code")})
        manifest["artifacts"] = [{"path": str(p.relative_to(self.path)), "sha256": file_hash(p), "size_bytes": p.stat().st_size} for p in artifact_paths]
        atomic_json(self.path / "run_manifest.json", manifest)
        if not failed:
            # The pointer is updated only after all artifacts and the completion manifest exist.
            atomic_json(self.root / "latest_successful.json", {"run_id": self.run_id, "completed_at": utc_now()})
        return manifest

    def close(self):
        self.logger.removeHandler(self.log_handler)
        self.logger.removeHandler(self.collector)
        self.log_handler.close()
        self.logger.setLevel(self.old_level)


def process_alive(pid):
    """Conservative local liveness check; never signal/terminate a Windows process."""
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
        kernel.OpenProcess.restype = wintypes.HANDLE
        kernel.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        handle = kernel.OpenProcess(0x00100000, False, pid)  # SYNCHRONIZE only
        if not handle:
            return ctypes.get_last_error() != 87  # INVALID_PARAMETER: PID does not exist
        try:
            return kernel.WaitForSingleObject(handle, 0) != 0
        finally:
            kernel.CloseHandle(handle)
    try:
        os.kill(pid, 0)
        return True
    except ProcessLookupError:
        return False
    except PermissionError:
        return True


def recover_interrupted_runs(root):
    """Finalize abandoned local RUNNING records. Active/foreign-host runs are untouched."""
    recovered = []
    for summary_path in Path(root).glob("run_*/run_summary.json"):
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
        if summary.get("execution_status") != "RUNNING":
            continue
        if summary.get("owner_host") != socket.gethostname() or process_alive(summary.get("owner_pid", os.getpid())):
            continue
        summary.update(execution_status="INTERRUPTED", decision_status="BLOCKED", exit_code=130,
                       finished_at=utc_now(), duration_seconds=None,
                       error={"type": "AbandonedRun", "message": "Owning process exited before finalization; rerun into a new directory"})
        for stage in summary["stages"]:
            if stage["status"] == "RUNNING":
                stage["status"] = "INTERRUPTED"
        summary["artifacts"] = [str(p.relative_to(summary_path.parent)) for p in summary_path.parent.rglob("*")
                                if p.is_file() and p.parent.name in {"artifacts", "artifacts.partial"}]
        atomic_json(summary_path, summary)
        report = summary_path.with_name("run_summary.md")
        report.write_text(f"# Skill ID Run Summary\n\nRun: {summary['run_id']}\n\nExecution: INTERRUPTED\n\nDecision status: BLOCKED\n\nThe owning process exited before finalization. Partial artifacts are not complete results.\nRerun the original command; a new run directory will be created.\n\nSee run_summary.json for stage and source details.\n", encoding="utf-8")
        manifest = {"run_id": summary["run_id"], "execution_status": "INTERRUPTED", "exit_code": 130,
                    "artifacts": [{"path": p.name, "sha256": file_hash(p), "size_bytes": p.stat().st_size} for p in (summary_path, report)]}
        atomic_json(summary_path.with_name("run_manifest.json"), manifest)
        recovered.append(summary["run_id"])
    return recovered
