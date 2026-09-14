"""Portable Python Excel authoring for the requested reusable pipeline."""
from pathlib import Path
from math import ceil
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


def style_table(sheet, rows, columns, number=1):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="254665")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 42
    for c in range(1, columns+1):
        values = [str(sheet.cell(r,c).value or "") for r in range(1,min(rows,100)+1)]
        sheet.column_dimensions[get_column_letter(c)].width = min(max(max(map(len, values),default=12)+2,16),42)
    if rows>1:
        table = Table(displayName=f"SkillTable{number}",ref=f"A1:{get_column_letter(columns)}{rows}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
        sheet.add_table(table)


def write_result_workbook(path, tables):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index,(name,frame) in enumerate(tables.items(),1):
        sheet = workbook.create_sheet(name[:31])
        if not len(frame.columns):
            frame = pd.DataFrame({"Status":["No records"]})
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False,name=None):
            safe=[]
            for value in row:
                if pd.isna(value):
                    value=None
                elif isinstance(value,pd.Timestamp):
                    value=value.to_pydatetime()
                elif hasattr(value,"item"):
                    value=value.item()
                if isinstance(value,str) and value.startswith(("=","+","-","@")):
                    value="'"+value
                safe.append(value)
            sheet.append(safe)
        style_table(sheet,sheet.max_row,sheet.max_column,index)
        sheet.sheet_format.defaultRowHeight = 20.0
        cell_font = Font(name="Arial", size=11)
        cell_align = Alignment(wrap_text=True, vertical="center")
        if sheet.max_row <= 2000:
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = cell_font
                    cell.alignment = cell_align
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000"
        else:
            float_cols = [i + 1 for i, dtype in enumerate(frame.dtypes) if pd.api.types.is_float_dtype(dtype)]
            for col_idx in float_cols:
                for r in range(2, sheet.max_row + 1):
                    sheet.cell(row=r, column=col_idx).number_format = "0.000"
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    workbook.save(path)
